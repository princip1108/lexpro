from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SOURCE = PROJECT_ROOT / "docs" / "LexPro多模态实体识别模型部署说明书.md"
TARGET = Path(os.environ.get(
    "LEXPRO_DOCX_TARGET",
    PROJECT_ROOT / "docs" / "LexPro多模态实体识别模型部署说明书.docx",
))


def set_run_font(run, latin="Calibri", east_asia="宋体", size=11, bold=None, color=None):
    run.font.name = latin
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(paragraph, fill="F2F4F7"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color="666666")


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


def new_numbering_id(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_ids, default=0) + 1
    style_num_id = doc.styles["List Number"]._element.pPr.numPr.numId.val
    base_num = next(node for node in numbering.findall(qn("w:num")) if int(node.get(qn("w:numId"))) == style_num_id)
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId"); abstract.set(qn("w:val"), abstract_id); num.append(abstract)
    override = OxmlElement("w:lvlOverride"); override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride"); start.set(qn("w:val"), "1")
    override.append(start); num.append(override); numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    num_pr.get_or_add_ilvl().set(qn("w:val"), "0")
    num_pr.get_or_add_numId().set(qn("w:val"), str(num_id))


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"; style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def build():
    doc = Document()
    configure_styles(doc)

    header = doc.sections[0].header.paragraphs[0]
    header.text = "LexPro 多模态询问笔录实体识别系统部署说明书"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        set_run_font(run, east_asia="微软雅黑", size=9, color="666666")
    add_page_number(doc.sections[0].footer.paragraphs[0])

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(150)
    run = cover.add_run("LexPro 多模态询问笔录\n实体识别系统部署说明书")
    set_run_font(run, east_asia="微软雅黑", size=24, bold=True, color="1F4D78")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(42)
    run = sub.add_run("版本 1.0")
    set_run_font(run, east_asia="微软雅黑", size=13, color="555555")
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(110)
    run = date.add_run("2026年7月")
    set_run_font(run, east_asia="微软雅黑", size=14, color="333333")
    doc.add_page_break()

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    skipped_title = False
    active_numbering_id = None
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            active_numbering_id = None
            if not in_code:
                in_code = True; code_lines = []
            else:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.12)
                paragraph.paragraph_format.right_indent = Inches(0.12)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(6)
                shade(paragraph)
                run = paragraph.add_run("\n".join(code_lines))
                set_run_font(run, latin="Consolas", east_asia="等线", size=8.5, color="333333")
                in_code = False
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("# ") and not skipped_title:
            skipped_title = True
            continue
        if line.startswith("**版本") or line.startswith("**日期"):
            continue
        if line.startswith("## "):
            active_numbering_id = None
            doc.add_paragraph(clean_inline(line[3:]), style="Heading 1")
        elif line.startswith("### "):
            active_numbering_id = None
            doc.add_paragraph(clean_inline(line[4:]), style="Heading 2")
        elif line.startswith("#### "):
            active_numbering_id = None
            doc.add_paragraph(clean_inline(line[5:]), style="Heading 3")
        elif re.match(r"^\d+\.\s", line):
            text = re.sub(r"^\d+\.\s*", "", line)
            if active_numbering_id is None:
                active_numbering_id = new_numbering_id(doc)
            paragraph = doc.add_paragraph(clean_inline(text), style="List Number")
            apply_numbering(paragraph, active_numbering_id)
        elif line.startswith("- "):
            active_numbering_id = None
            doc.add_paragraph(clean_inline(line[2:]), style="List Bullet")
        elif line.strip():
            active_numbering_id = None
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(clean_inline(line))
            set_run_font(run)

    props = doc.core_properties
    props.title = "LexPro 多模态询问笔录实体识别系统部署说明书"
    props.subject = "本地化部署、服务启动、推理与验收"
    props.author = "LexPro 项目组"
    props.keywords = "LexPro,MinerU,vLLM,多模态实体识别"
    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    build()
